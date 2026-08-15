from __future__ import annotations

from pathlib import Path

from pypdf import PdfReader
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


def load_document_text(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="replace").strip()

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        return text.strip()

    doc = Document(str(path))
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        if rows:
            blocks.append("\n".join(rows))
    return "\n\n".join(blocks).strip()
